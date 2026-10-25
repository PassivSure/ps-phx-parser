"""Heat-pump capacity and COP, read off the HP sheet's test-point tables.

The workbooks here are laid out like a real v10.6 HP sheet: the selection and
active-flag names at their real cells, and each section's test-point table at
its real rows, so a row offset that drifts fails here rather than only on the
real-workbook suite (which skips when the files are absent).
"""

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from app.equipment import read_equipment

# name -> cell, as PHPP v10.6 defines them (both editions).
NAMED_CELLS = {
    "WP_Heizung_ankreuzen": "I22",
    "WP_Heizung_Auswahl_Waermepumpe": "K23",
    "WP_LuftLuft_ankreuzen": "I43",
    "WP_Heizungssystem_Auswahl_Luft_Luft_WP": "K45",
    "WP_WWseparat_ankreuzen": "I70",
    "WP_Warmwassersystem_Auswahl_WP": "K71",
}

# section -> first test-point row
TABLE_ROWS = {"water": 27, "air": 50, "dhw": 84}


def _hp_wb(*, sheet="HP SI", values=None, tables=None, unit="kW", label="Test point 1"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for name, cell in NAMED_CELLS.items():
        wb.defined_names.add(DefinedName(name, attr_text=f"'{sheet}'!${cell[0]}${cell[1:]}"))
    for name, value in (values or {}).items():
        ws[NAMED_CELLS[name]] = value
    for section, points in (tables or {}).items():
        start = TABLE_ROWS[section]
        ws[f"AF{start - 1}"] = unit
        for i, (source, sink, capacity, cop) in enumerate(points):
            row = start + i
            ws[f"AC{row}"] = label if i == 0 else f"Test point {i + 1}"
            ws[f"AD{row}"], ws[f"AE{row}"], ws[f"AF{row}"], ws[f"AG{row}"] = source, sink, capacity, cop
    return wb


AIR = [(10, 21, 18.36, 4.8), (7, 21, 17, 4.7), (2, 21, 12.24, 3.8)]
DHW = [(2, 50, 5.04, 2.2), (7, 50, 6, 2.5), (15, 50, 7.44, 3.0)]
WATER = [(-7, 35, 2.88, 2.7), (15, 35, 5, 4.3), (15, 50, 4.96, 3.3)]


def _of(items, equipment_type):
    return [i for i in items if i["equipment_type"] == equipment_type]


class TestAirToAir:
    def test_reads_the_largest_capacity_and_the_cop_at_that_point(self):
        wb = _hp_wb(values={"WP_LuftLuft_ankreuzen": "x",
                            "WP_Heizungssystem_Auswahl_Luft_Luft_WP": "99ud-Standard air-to-air heat pump"},
                    tables={"air": AIR})
        [hp] = _of(read_equipment(wb, "EN_10_6IP"), "heat_pump")
        assert hp["name"] == "Standard air-to-air heat pump"
        assert (hp["capacity"], hp["capacity_unit"]) == (18.36, "kW")
        assert (hp["efficiency_value"], hp["efficiency_type"]) == (4.8, "COP")
        assert hp["source"].startswith("WP_Heizungssystem_Auswahl_Luft_Luft_WP; test point 1")

    def test_a_selection_whose_box_is_not_ticked_is_not_emitted(self):
        wb = _hp_wb(values={"WP_Heizungssystem_Auswahl_Luft_Luft_WP": "99ud-Standard air-to-air heat pump"},
                    tables={"air": AIR})
        assert _of(read_equipment(wb, "EN_10_6IP"), "heat_pump") == []

    def test_skips_zero_and_error_cells_when_picking_the_point(self):
        points = [(0, 60, 0, 2.18), (10, 21, "#VALUE!", 4.9), (20, 60, 1.6, 2.84)]
        wb = _hp_wb(values={"WP_LuftLuft_ankreuzen": "x",
                            "WP_Heizungssystem_Auswahl_Luft_Luft_WP": "01ud-Unit"},
                    tables={"air": points})
        [hp] = _of(read_equipment(wb, "EN_10_6IP"), "heat_pump")
        assert (hp["capacity"], hp["efficiency_value"]) == (1.6, 2.84)
        assert "test point 3" in hp["source"]

    def test_an_empty_table_keeps_the_unit_with_no_figures(self):
        wb = _hp_wb(values={"WP_LuftLuft_ankreuzen": "x",
                            "WP_Heizungssystem_Auswahl_Luft_Luft_WP": "01ud-Unit"})
        [hp] = _of(read_equipment(wb, "EN_10_6IP"), "heat_pump")
        assert hp["capacity"] is None and hp["efficiency_value"] is None
        assert hp["capacity_unit"] is None and hp["efficiency_type"] is None


class TestLabelCheck:
    """The rows are fixed, so a layout that moved must read nothing rather than
    the wrong row."""

    def test_a_table_not_labelled_test_point_1_reads_nothing(self):
        wb = _hp_wb(values={"WP_LuftLuft_ankreuzen": "x",
                            "WP_Heizungssystem_Auswahl_Luft_Luft_WP": "01ud-Unit"},
                    tables={"air": AIR}, label="Something else")
        [hp] = _of(read_equipment(wb, "EN_10_6IP"), "heat_pump")
        assert hp["capacity"] is None
        assert "test points not found" in hp["source"]

    def test_a_table_not_in_kw_reads_nothing(self):
        wb = _hp_wb(values={"WP_LuftLuft_ankreuzen": "x",
                            "WP_Heizungssystem_Auswahl_Luft_Luft_WP": "01ud-Unit"},
                    tables={"air": AIR}, unit="kBTU/hr")
        [hp] = _of(read_equipment(wb, "EN_10_6IP"), "heat_pump")
        assert hp["capacity"] is None
        assert "test points not found" in hp["source"]


class TestDhw:
    def test_reads_the_separate_dhw_heat_pump(self):
        wb = _hp_wb(values={"WP_WWseparat_ankreuzen": "x",
                            "WP_Warmwassersystem_Auswahl_WP": "95ud-Standard DHW HP"},
                    tables={"dhw": DHW})
        [dhw] = _of(read_equipment(wb, "EN_10_6IP"), "dhw_heat_pump")
        assert (dhw["capacity"], dhw["efficiency_value"]) == (7.44, 3.0)

    def test_a_dhw_selection_whose_box_is_not_ticked_is_not_emitted(self):
        wb = _hp_wb(values={"WP_Warmwassersystem_Auswahl_WP": "95ud-Standard DHW HP"},
                    tables={"dhw": DHW})
        assert _of(read_equipment(wb, "EN_10_6IP"), "dhw_heat_pump") == []


class TestWaterBased:
    def test_reads_a_ticked_water_based_heat_pump(self):
        wb = _hp_wb(sheet="HP", values={"WP_Heizung_ankreuzen": "x",
                                        "WP_Heizung_Auswahl_Waermepumpe": "96ud-Standard air/water heat pump"},
                    tables={"water": WATER})
        [hp] = _of(read_equipment(wb, "EN_10_6"), "heat_pump")
        assert hp["name"] == "Standard air/water heat pump"
        assert (hp["capacity"], hp["efficiency_value"]) == (5, 4.3)
        assert hp["source"].startswith("WP_Heizung_Auswahl_Waermepumpe")

    def test_an_unticked_water_based_selection_is_not_emitted(self):
        wb = _hp_wb(values={"WP_Heizung_Auswahl_Waermepumpe": "96ud-Standard air/water heat pump"},
                    tables={"water": WATER})
        assert read_equipment(wb, "EN_10_6IP") == []


def test_reads_the_sheet_the_selection_name_points_at():
    """An IP workbook's names point at 'HP SI'; the IP pane beside it shows the
    COP multiplied by 3.412. Reading anything but the named sheet would bring
    that back."""
    wb = _hp_wb(values={"WP_LuftLuft_ankreuzen": "x",
                        "WP_Heizungssystem_Auswahl_Luft_Luft_WP": "01ud-Unit"},
                tables={"air": AIR})
    ip = wb.create_sheet("HP")
    ip["AC50"], ip["AF50"], ip["AG50"] = "Test point 1", 62.65, 16.38
    [hp] = _of(read_equipment(wb, "EN_10_6IP"), "heat_pump")
    assert hp["efficiency_value"] == 4.8
