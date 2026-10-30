"""Windows from the Windows sheet (PAS-80, P2.3.c).

Synthetic sheets pin the reading rules; the real workbooks (skipped when the
files are absent) pin that what is read adds up to each sheet's own totals,
Windows!V16 (area) and W16 (area-weighted Uw installed) on v10, read by hand.
"""

import pathlib

import pytest
from openpyxl import Workbook, load_workbook

from app.envelope import read_windows
from app.parser import load_shape, parse_workbook

DOWNLOADS = pathlib.Path.home() / "Downloads"
FT2_PER_M2 = 10.763910416709722
IP_U_PER_SI_U = 0.17611  # Btu/(h·ft²·°F) per W/(m²K)


def _windows_sheet(rows, shape_name="EN_10_6IP", header_row=21, g_label="g-value\nglazing"):
    """A Windows sheet with the shape's locators, its header labels, and +rows+
    (dicts of quantity, description, glazing, frame, area, u, g) from header+2."""
    shape = load_shape(shape_name)
    w = shape.WINDOWS
    cols = w.window_rows.inputs
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(w.name)
    ws[f"{w.window_rows.locator_col_header}3"] = w.window_rows.locator_string_header
    ws[f"{w.window_rows.locator_col_entry}{header_row}"] = "Quan-\ntity"
    ws[f"AS{header_row}"] = g_label
    r = header_row + 2
    fields = {
        "quantity": "quantity",
        "description": "description",
        "glazing_id": "glazing",
        "frame_id": "frame",
        "window_area": "area",
        "u_w_installed": "u",
        "u_w": "u_w",
    }
    for row in rows:
        for name, key in fields.items():
            column = getattr(cols, name).column
            if column:  # v9.7 maps no frame column
                ws[f"{column}{r}"] = row.get(key, 1 if key == "quantity" else None)
        ws[f"AS{r}"] = row.get("g")
        r += 1
    ws[f"{w.window_rows_end.locator_col_entry}{r + 3}"] = w.window_rows_end.locator_string_entry
    ws[f"{cols.window_area.column}{r + 5}"] = 999.0  # past the end marker: never read
    return wb, shape


class TestReadWindows:
    def test_reads_each_window_row_with_installed_uw_g_and_ids(self):
        wb, shape = _windows_sheet(
            [
                {
                    "description": "W1",
                    "glazing": "01ud-Triple Kr",
                    "frame": "02ud-Operable",
                    "area": 14.5,
                    "u": 0.140,
                    "u_w": 0.120,
                    "g": 0.5,
                    "quantity": 2,
                },
            ]
        )

        [w] = read_windows(wb, shape)

        assert w == {
            "component": "window",
            "label": "W1",
            "quantity": 2,
            "area_ft2": 14.5,
            "u_value_Btuh_ft2F": 0.140,  # Uw installed (BA), never Uw (AZ)
            "shgc": 0.5,
            "glazing_id": "01ud",
            "frame_id": "02ud",
            "source_area": "Windows!AW23",
            "source_u_value": "Windows!BA23",
            "source_shgc": "Windows!AS23",
        }

    def test_skips_rows_with_no_area_and_stops_at_the_end_marker(self):
        wb, shape = _windows_sheet(
            [
                {"description": "W1", "area": 10.0, "u": 0.14, "g": 0.5},
                {"description": None, "area": 0, "u": 0, "g": 0},  # PHPP's empty row
                {"description": "W2", "area": 5.0, "u": 0.15, "g": 0.4},
            ]
        )

        assert [w["label"] for w in read_windows(wb, shape)] == ["W1", "W2"]

    def test_converts_an_si_sheet_to_the_imperial_keys(self):
        wb, shape = _windows_sheet(
            [{"description": "W1", "area": 2.0, "u": 0.8, "g": 0.5}], shape_name="EN_10_6"
        )

        [w] = read_windows(wb, shape)

        assert w["area_ft2"] == pytest.approx(2.0 * FT2_PER_M2)
        assert w["u_value_Btuh_ft2F"] == pytest.approx(0.8 * IP_U_PER_SI_U, rel=1e-4)

    def test_reads_v9_7s_header_a_row_lower_and_its_shgc_label(self):
        wb, shape = _windows_sheet(
            [{"description": "d02", "area": 38.25, "u": 0.178, "g": 0.44}],
            shape_name="EN_9_7IP",
            header_row=22,
            g_label="SHGC",
        )

        [w] = read_windows(wb, shape)

        assert (w["area_ft2"], w["shgc"]) == (38.25, 0.44)

    def test_leaves_shgc_null_when_no_g_column_is_labelled(self):
        wb, shape = _windows_sheet(
            [{"description": "W1", "area": 1.0, "u": 0.1, "g": 0.5}], g_label="something else"
        )

        [w] = read_windows(wb, shape)

        assert w["shgc"] is None
        assert "source_shgc" not in w

    def test_reads_only_inside_the_bounded_sheet_when_there_is_no_end_marker(self):
        """A sheet with no end marker reads to the bound, never to a declared max_row."""
        shape = load_shape("EN_10_6IP")
        w = shape.WINDOWS
        cols = w.window_rows.inputs
        wb = Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet(w.name)
        ws[f"{w.window_rows.locator_col_entry}21"] = "Quan-\ntity"
        for r, label in ((23, "W1"), (401, "past the bound")):
            ws[f"{cols.description.column}{r}"] = label
            ws[f"{cols.window_area.column}{r}"] = 10.0
            ws[f"{cols.u_w_installed.column}{r}"] = 0.14
        ws.cell(row=1_048_576, column=1, value="far")

        assert [x["label"] for x in read_windows(wb, shape)] == ["W1"]

    def test_is_empty_without_a_windows_sheet(self):
        wb = Workbook()
        assert read_windows(wb, load_shape("EN_10_6IP")) == []


# file, shape. The totals checked are read off each sheet's own V16 / W16.
REAL = {
    "phi_10_6_ip": ("PHPP_V10.6_IP_Example.xlsx", "EN_10_6IP"),
    "donovan_10_6_ip": ("260309 PHPP 6840 E ^th Ave Pkwy v 10.6.xlsx", "EN_10_6IP"),
    "easyph_si": ("PHPP_EN_V10.6_easyPH_Example.xlsx", "EN_10_6"),
}


@pytest.mark.parametrize("case", sorted(REAL))
def test_real_windows_add_up_to_the_sheets_totals(case):
    filename, shape_name = REAL[case]
    path = DOWNLOADS / filename
    if not path.exists():
        pytest.skip(f"{filename} not in ~/Downloads")

    sheet = load_workbook(path, data_only=True, read_only=True)["Windows"]
    total_area, total_u = sheet["V16"].value, sheet["W16"].value
    si = shape_name == "EN_10_6"
    total_area_ft2 = total_area * (FT2_PER_M2 if si else 1)
    total_u_ip = total_u * (IP_U_PER_SI_U if si else 1)

    windows = [
        c
        for c in parse_workbook(path.read_bytes(), shape_name)["envelope"]["components"]
        if c["component"] == "window"
    ]

    assert len(windows) >= 3
    area = sum(w["area_ft2"] for w in windows)
    assert area == pytest.approx(total_area_ft2, rel=1e-6)
    assert sum(w["area_ft2"] * w["u_value_Btuh_ft2F"] for w in windows) / area == pytest.approx(
        total_u_ip, rel=1e-3
    )
    assert all(w["glazing_id"] and 0 < w["shgc"] < 1 for w in windows)
