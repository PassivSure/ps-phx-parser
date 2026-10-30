"""app.assemblies on synthetic workbooks laid out like a real v10.6 file: the
U-values blocks and the Components assemblies table at their real cells, found
through the same defined names."""

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from app.assemblies import read_assemblies, split_assembly_ref

U_SHEET = "U-values SI"
C_SHEET = "Components SI"


def _name(wb, name, ref):
    wb.defined_names.add(DefinedName(name, attr_text=ref))


def _block(wb, n, *, assembly_id, name, layers, section_2=0.0, u=None,
           header="Description of building assembly"):
    """Block n at its v10.6 rows: name row 8 + 21(n-1), layers 5 rows below."""
    ws = wb[U_SHEET]
    r = 8 + 21 * (n - 1)
    ws[f"L{r - 1}"], ws[f"Q{r - 1}"] = header, "Assembly no."
    ws[f"L{r}"], ws[f"Q{r}"] = name, assembly_id
    for i, layer in enumerate(layers):
        row = r + 5 + i
        for col, value in zip("LMNOPQR", layer, strict=False):
            ws[f"{col}{row}"] = value
    ws[f"O{r + 13}"] = section_2
    ws[f"Q{r + 17}"], ws[f"R{r + 17}"] = "U-value [W/(m²K)]:", u
    _name(wb, f"U_Werte_Bauteil_{n}_Bezeichnung", f"'{U_SHEET}'!$L${r}:$P${r}")
    _name(wb, f"U_Werte_Bauteil_{n}_Teilflaechen", f"'{U_SHEET}'!$L${r + 5}:$Q${r + 12}")
    _name(wb, f"U_Werte_Bauteil_{n}_Summe_Breite_Dicke", f"'{U_SHEET}'!$R${r + 5}:$R${r + 12}")
    _name(wb, f"U_Werte_Bauteil_{n}_Flaechenanteil_2", f"'{U_SHEET}'!$O${r + 13}")
    _name(wb, f"U_Werte_Bauteil_{n}_Flaechenanteil_3", f"'{U_SHEET}'!$Q${r + 13}")


def _table(wb, rows, unit="W/(m²K)"):
    """The Components building-assemblies table: heading E8, labels row 11,
    units row 12, data from row 13."""
    ws = wb[C_SHEET]
    ws["G8"] = "Building assemblies (U-values)"
    labels = ("ID", "Building system", "Building assembly", "\nTotal\nthickness\n", "U-value")
    for col, label in zip("GHIJK", labels, strict=True):
        ws[f"{col}11"] = label
    ws["J12"], ws["K12"] = "m", unit
    for i, (assembly_id, name, thickness_m, u) in enumerate(rows):
        r = 13 + i
        ws[f"G{r}"], ws[f"I{r}"], ws[f"J{r}"], ws[f"K{r}"] = assembly_id, name, thickness_m, u
    _name(wb, "PHPP_Komponenten_Ueberschrift_Bauteilaufbauten", f"'{C_SHEET}'!$E$8:$L$8")


def _wb():
    wb = Workbook()
    wb.active.title = U_SHEET
    wb.create_sheet(C_SHEET)
    return wb


def _surface(raw, area, component="wall", group=8):
    return {"assembly_id_raw": raw, "area_ft2": area, "component": component, "group": group}


WALL_LAYERS = [
    ("Concrete", 7.21, "Concrete", 7.21, None, 0, 203.2),
    ("Mineral Wool Batt", 0.041, '16" O.C.- 2x6', 0.152, None, 0, 139.7),
    ("GWB", 0.16, None, 0, None, 0, 12.7),
]


def _donovan_like():
    wb = _wb()
    _block(wb, 1, assembly_id="02ud", name="Wall 1- Basement", layers=WALL_LAYERS,
           section_2=0.09375, u=0.1542)
    _block(wb, 2, assembly_id="07ud", name=" ", layers=[], u=3.84)  # unused block
    _table(wb, [("01ud", "Slab", 0.3556, 0.1485), ("02ud", "Wall 1- Basement", 0.3556, 0.1542),
                ("83ud", None, 0.46, 0.15)])
    return wb


def test_reads_the_u_value_off_the_components_table_and_the_layers_off_the_block():
    [a] = read_assemblies(_donovan_like(), [_surface("02ud-Wall 1- Basement", 100.0, group=9)])
    assert a["id"] == "02ud" and a["name"] == "Wall 1- Basement"
    assert a["u_value_w_m2k"] == 0.1542 and a["source_u_value"] == "Components SI!K14"
    assert a["section_2_fraction"] == 0.09375
    assert [layer["material"] for layer in a["layers"]] == ["Concrete", "Mineral Wool Batt", "GWB"]
    stud = a["layers"][1]
    assert stud == {
        "thickness_mm": 139.7, "material": "Mineral Wool Batt", "conductivity_w_mk": 0.041,
        "section_2_material": '16" O.C.- 2x6', "section_2_conductivity_w_mk": 0.152,
        "section_3_material": None, "section_3_conductivity_w_mk": None,
    }
    assert a["total_thickness_mm"] == 203.2 + 139.7 + 12.7
    assert a["source_layers"] == "U-values SI!L13:Q20"


def test_a_builtin_assembly_has_its_u_and_no_layers():
    """83ud has no U-values block; its name is blank on the metric pane, so it
    comes from the Areas reference."""
    [a] = read_assemblies(_donovan_like(), [_surface("83ud-PH External wall", 24.0)])
    assert (a["name"], a["u_value_w_m2k"], a["layers"]) == ("PH External wall", 0.15, [])
    assert a["total_thickness_mm"] == 460.0


def test_sums_area_and_takes_the_type_with_the_most_area():
    surfaces = [_surface("02ud-Wall 1- Basement", 300.0),
                _surface("02ud-Wall 1- Basement", 20.0, None, 14)]
    [a] = read_assemblies(_donovan_like(), surfaces)
    assert (a["area_ft2"], a["surface_count"]) == (320.0, 2)
    assert (a["component"], a["groups"]) == ("wall", [8, 14])


def test_never_invents_a_type():
    [a] = read_assemblies(_donovan_like(), [_surface("83ud-Wall to zone X", 12.0, None, 14)])
    assert a["component"] is None and a["groups"] == [14]


def test_an_area_missing_on_one_surface_leaves_the_total_unknown():
    surfaces = [_surface("02ud-Wall 1- Basement", 300.0), _surface("02ud-Wall 1- Basement", None)]
    [a] = read_assemblies(_donovan_like(), surfaces)
    assert a["area_ft2"] is None


def test_an_unknown_id_is_emitted_with_nothing_read():
    [a] = read_assemblies(_donovan_like(), [_surface("42ud-Mystery", 5.0)])
    assert (a["id"], a["name"], a["u_value_w_m2k"], a["layers"]) == ("42ud", "Mystery", None, [])


def test_an_unused_block_is_not_read():
    """Block 2 has an id and a U but a blank name, as PHPP leaves unused ones."""
    [a] = read_assemblies(_donovan_like(), [_surface("07ud-", 5.0)])
    assert a["layers"] == [] and a["name"] is None and a["source_layers"] is None


def test_a_block_whose_header_moved_reads_no_layers():
    wb = _wb()
    _block(wb, 1, assembly_id="02ud", name="Wall", layers=WALL_LAYERS, header="Something else")
    _table(wb, [("02ud", "Wall", 0.36, 0.15)])
    [a] = read_assemblies(wb, [_surface("02ud-Wall", 1.0)])
    assert a["layers"] == [] and a["u_value_w_m2k"] == 0.15


def test_a_table_not_in_w_per_m2k_reads_no_u():
    """The IP pane's column is an R in hr·ft²·°F/Btu."""
    wb = _wb()
    _table(wb, [("02ud", "Wall", 0.36, 37.86)], unit="hr.ft2.F/BTU")
    [a] = read_assemblies(wb, [_surface("02ud-Wall", 1.0)])
    assert a["u_value_w_m2k"] is None


def test_no_surface_names_an_assembly():
    assert read_assemblies(_donovan_like(), [_surface(None, 5.0)]) == []


def test_split_assembly_ref():
    assert split_assembly_ref("01ud-Wall 1- Basement") == ("01ud", "Wall 1- Basement")
    assert split_assembly_ref("1363vs03-Brink - Flair") == ("1363vs03", "Brink - Flair")
    assert split_assembly_ref("Wall") == (None, None)
    assert split_assembly_ref(None) == (None, None)
